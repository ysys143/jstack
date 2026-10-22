#!/usr/bin/env python3
"""Assemble a slide manuscript into PPTX, optionally with PDF.

Usage:
  python3 assemble.py deck.md
  python3 assemble.py deck.md -o output/my_deck.pptx
  python3 assemble.py slide_plan.json -o output/my_deck.pptx

Markdown is a first-class input because deck retains it as source (Phase 4).
The legacy md-to-presentation script accepted only slide_plan.json, requiring
one finished image per slide and therefore an image-generation pipeline.
Keep JSON for cases where image paths already exist.

Exit codes: 0 success; 1 invalid/missing input, zero slides, or unknown format;
2 missing dependencies, reported with requirements.txt installation guidance.
Dependency failures need a different user remedy from input errors and should
not require reading a traceback.

Set CJK fonts in title/body runs including the East Asian a:ea slot, following
references/styles.md. Latin-only declarations leave Korean to machine-dependent
PowerPoint fallback. Use slide line height 1.2~1.35; web values 1.6~1.8 leave room
for only a few lines per slide.
"""

import argparse
import importlib
import json
import re
import sys
from pathlib import Path

# Slide dimensions: 16:9 widescreen; 914400 EMU per inch.
SLIDE_W_EMU = 12192000  # 33.87 cm
SLIDE_H_EMU = 6858000   # 19.05 cm

# CJK font stack, ordered according to references/styles.md.
CJK_FONT = "Pretendard"
CJK_FALLBACK = "Noto Sans KR"

# Slide line spacing within the 1.2~1.35 contract in references/styles.md.
LINE_SPACING = 1.25

# Minimum body size in points; styles.md shared rule 1.
BODY_PT_MIN = 14

DEPS = {
    "pptx": "python-pptx",
    "PIL": "pillow",
    "fpdf": "fpdf2",
}


def _need(mod: str):
    """Import a dependency lazily; absence exits 2 with installation guidance."""
    try:
        return importlib.import_module(mod)
    except ImportError:
        pkg = DEPS.get(mod, mod)
        print(
            f"Error: '{pkg}' is missing; this step cannot continue.\n"
            f"  Install: pip install -r requirements.txt   (deck skill directory)\n"
            f"  Or: pip install {pkg}\n"
            f"  requirements.txt lists all dependencies required by this script.",
            file=sys.stderr,
        )
        raise SystemExit(2)


# ── Read manuscript ────────────────────────────────────────────────────────────────

SEP = re.compile(r"^-{3,}\s*$")
HEAD = re.compile(r"^(#{1,3})\s+(.*\S)\s*$")
BULLET = re.compile(r"^\s*[-*+]\s+(.*\S)\s*$")


def parse_markdown(text: str) -> list[dict]:
    """Split Markdown into slides at --- lines (the marp convention).

    Remove leading YAML frontmatter, which is not a slide.
    """
    lines = text.splitlines()
    if lines and SEP.match(lines[0]):
        for i in range(1, len(lines)):
            if SEP.match(lines[i]):
                lines = lines[i + 1:]
                break

    chunks, cur = [], []
    for line in lines:
        if SEP.match(line):
            chunks.append(cur)
            cur = []
        else:
            cur.append(line)
    chunks.append(cur)

    slides = []
    for chunk in chunks:
        title, bullets, body = "", [], []
        for line in chunk:
            m = HEAD.match(line)
            if m:
                if not title:
                    title = m.group(2)
                else:
                    body.append(m.group(2))
                continue
            m = BULLET.match(line)
            if m:
                bullets.append(m.group(1))
                continue
            if line.strip():
                body.append(line.strip())
        if not (title or bullets or body):
            continue
        slides.append({
            "slide_num": len(slides) + 1,
            "title": title,
            "key_points": bullets,
            "body": body,
            "image_path": "",
        })
    return slides


def parse_plan(text: str) -> list[dict]:
    """Read legacy slide_plan.json, which assumes one finished image per slide."""
    plan = json.loads(text)
    slides = plan.get("slides") or []
    out = []
    for s in slides:
        out.append({
            "slide_num": s.get("slide_num", len(out) + 1),
            "title": s.get("title", ""),
            "key_points": s.get("key_points", []),
            "body": [],
            "image_path": s.get("image_path", ""),
            "speaker_notes": s.get("speaker_notes", ""),
        })
    return sorted(out, key=lambda s: s["slide_num"])


def load(path: Path) -> list[dict]:
    text = path.read_text(encoding="utf-8")
    if path.suffix.lower() == ".json":
        return parse_plan(text)
    if path.suffix.lower() in (".md", ".markdown", ".txt"):
        return parse_markdown(text)
    raise SystemExit(f"Error: unknown input format ({path.suffix}); use .md or .json")


# ── Typesetting ─────────────────────────────────────────────────────────────────────


def _set_cjk(run, size_pt, bold=False):
    """Set run fonts including the East Asian slot.

    run.font.name sets only a:latin. Korean uses a:ea; omitting it invokes
    viewer defaults, the original defect this skill's contract prevents.
    """
    from pptx.oxml.ns import qn
    from pptx.util import Pt

    run.font.size = Pt(size_pt)
    run.font.bold = bold
    run.font.name = CJK_FONT          # a:latin: Latin slot.
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):      # East Asian and complex-script slots.
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {})
            rPr.append(el)
        el.set("typeface", CJK_FONT)
    return run


def _fit_image(img_path: str):
    """Return centered placement (left, top, width, height) in EMU."""
    _need("PIL")
    from PIL import Image
    from pptx.util import Emu

    img = Image.open(img_path)
    ratio = img.width / img.height
    slide_ratio = SLIDE_W_EMU / SLIDE_H_EMU
    if ratio > slide_ratio:
        w = SLIDE_W_EMU
        h = int(SLIDE_W_EMU / ratio)
        left, top = 0, (SLIDE_H_EMU - h) // 2
    else:
        h = SLIDE_H_EMU
        w = int(SLIDE_H_EMU * ratio)
        left, top = (SLIDE_W_EMU - w) // 2, 0
    return Emu(left), Emu(top), Emu(w), Emu(h)


def build_pptx(slides: list[dict], output_path: str) -> str:
    _need("pptx")
    from pptx import Presentation
    from pptx.util import Emu

    prs = Presentation()
    prs.slide_width = Emu(SLIDE_W_EMU)
    prs.slide_height = Emu(SLIDE_H_EMU)
    blank = prs.slide_layouts[6]

    for s in slides:
        slide = prs.slides.add_slide(blank)
        img = s.get("image_path", "")
        if img and Path(img).exists():
            _need("PIL")
            left, top, w, h = _fit_image(img)
            slide.shapes.add_picture(img, left, top, w, h)
            _overlay(slide, s)
        else:
            _text_slide(slide, s)

        notes = s.get("speaker_notes", "")
        if notes:
            tf = slide.notes_slide.notes_text_frame
            if tf is not None:
                tf.text = notes
        print(f"  Slide {s['slide_num']:02d}: {s.get('title', '')[:50]}")

    pptx_path = output_path if output_path.endswith(".pptx") else output_path + ".pptx"
    Path(pptx_path).parent.mkdir(parents=True, exist_ok=True)
    prs.save(pptx_path)
    print(f"  PPTX: {pptx_path} ({Path(pptx_path).stat().st_size / 1024:.0f}KB)")
    return pptx_path


def _text_slide(slide, s: dict):
    """Create a text slide, the default for Markdown manuscripts."""
    from pptx.util import Emu

    top = 600000
    if s.get("title"):
        box = slide.shapes.add_textbox(
            Emu(700000), Emu(top), Emu(SLIDE_W_EMU - 1400000), Emu(1100000))
        tf = box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.line_spacing = LINE_SPACING
        run = p.add_run()
        run.text = s["title"]
        _set_cjk(run, 32, bold=True)
        top += 1300000

    lines = [f"• {b}" for b in s.get("key_points", [])] + list(s.get("body", []))
    if not lines:
        return
    box = slide.shapes.add_textbox(
        Emu(700000), Emu(top), Emu(SLIDE_W_EMU - 1400000),
        Emu(max(600000, SLIDE_H_EMU - top - 500000)))
    tf = box.text_frame
    tf.word_wrap = True
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = LINE_SPACING
        run = p.add_run()
        run.text = line
        _set_cjk(run, max(BODY_PT_MIN, 18))


def _overlay(slide, s: dict):
    """Overlay title/key points on an image slide for the JSON route."""
    from pptx.dml.color import RGBColor
    from pptx.util import Emu

    if not s.get("title"):
        return
    box = slide.shapes.add_textbox(
        Emu(457200), Emu(5600000), Emu(8229600), Emu(600000))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.line_spacing = LINE_SPACING
    run = p.add_run()
    run.text = s["title"]
    _set_cjk(run, 24, bold=True)
    run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)


def build_pdf(slides: list[dict], pptx_path: str) -> str | None:
    """Emit PDF only for image slides.

    Text slides need embedded Korean TTF fonts for fpdf2, but this skill does
    not distribute fonts. Missing glyphs would otherwise count as successful
    rendering. Use the marp route (SKILL.md Phase 5) for Markdown manuscript PDFs.
    """
    imgs = [s for s in slides if s.get("image_path") and Path(s["image_path"]).exists()]
    if not imgs:
        print("  PDF:  skipped; no image slides (use marp for PDF)")
        return None

    _need("fpdf")
    _need("PIL")
    from fpdf import FPDF
    from PIL import Image

    pdf_path = pptx_path.replace(".pptx", ".pdf")
    page_w = SLIDE_W_EMU / 914400 * 72
    page_h = SLIDE_H_EMU / 914400 * 72
    pdf = FPDF(orientation="L", unit="pt", format=(page_h, page_w))
    pdf.set_auto_page_break(False)
    for s in imgs:
        pdf.add_page()
        img = Image.open(s["image_path"])
        ratio = img.width / img.height
        if ratio > page_w / page_h:
            w, h = page_w, page_w / ratio
            x, y = 0.0, (page_h - h) / 2
        else:
            h, w = page_h, page_h * ratio
            x, y = (page_w - w) / 2, 0.0
        pdf.image(s["image_path"], x=x, y=y, w=w, h=h)
    pdf.output(pdf_path)
    print(f"  PDF:  {pdf_path} ({Path(pdf_path).stat().st_size / 1024:.0f}KB)")
    return pdf_path


def main():
    parser = argparse.ArgumentParser(
        description="Assemble a slide manuscript (.md) or slide_plan.json into PPTX")
    parser.add_argument("input", help="Manuscript .md or slide_plan.json")
    parser.add_argument("-o", "--output", default="workspace/deck/output/result.pptx",
                        help="Output PPTX path (default: workspace/deck/output/result.pptx)")
    parser.add_argument("--skip-missing", action="store_true",
                        help="Omit slides without images when assembling the JSON route")
    parser.add_argument("--no-pdf", action="store_true", help="Do not create PDF")
    args = parser.parse_args()

    src = Path(args.input)
    if not src.exists():
        print(f"Error: {args.input} does not exist", file=sys.stderr)
        return 1

    slides = load(src)
    if args.skip_missing:
        slides = [s for s in slides
                  if s.get("image_path") and Path(s["image_path"]).exists()]
    if not slides:
        print("Error: zero slides; check the manuscript", file=sys.stderr)
        return 1

    print(f"=== Assemble {len(slides)} slides ===")
    pptx_path = build_pptx(slides, args.output)
    if not args.no_pdf:
        build_pdf(slides, pptx_path)
    print("=== Complete ===")
    return 0


if __name__ == "__main__":
    sys.exit(main())
