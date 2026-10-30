#!/usr/bin/env python3
"""Extract PDF presentation pages and reassemble them into PPTX plus PDF.

Usage:
  python3 pdf2pptx.py input.pdf                  # Extract, upscale, assemble
  python3 pdf2pptx.py input.pdf --skip-upscale   # Keep original resolution
  python3 pdf2pptx.py input.pdf --resume         # Skip completed pages
  python3 pdf2pptx.py input.pdf -o out/deck.pptx

The legacy pdf-to-presentation script duplicated Gemini REST call_gemini and
model IDs also owned by the image skill. Updating one copy could silently diverge
from the other. Upscaling belongs to image: invoke its call_gemini and exit 2
when unavailable rather than copying the generator here.

Exit codes: 0 success; 1 missing PDF or zero pages; 2 missing dependencies,
reported with requirements.txt or image-skill guidance.

Recommend --skip-upscale first: it needs neither the image skill nor an API key,
original images are often large enough, and upscaling cannot preserve layout
with complete fidelity.
"""

import argparse
import importlib
import importlib.util
import json
import sys
import time
from pathlib import Path

SLIDE_W_EMU = 12192000  # 16:9 widescreen
SLIDE_H_EMU = 6858000

DEPS = {
    "fitz": "pymupdf",
    "PIL": "pillow",
    "pptx": "python-pptx",
    "fpdf": "fpdf2",
}

UPSCALE_PROMPT = (
    "Upscale this presentation slide to higher resolution. "
    "Preserve all text, layout, colors, and visual elements exactly. "
    "Only enhance resolution and sharpness."
)

WORKSPACE = Path("workspace")
EXTRACTED_DIR = WORKSPACE / "extracted"
UPSCALED_DIR = WORKSPACE / "upscaled"
OUTPUT_DIR = WORKSPACE / "output"
STATE_PATH = WORKSPACE / "state.json"

# Seconds between upscale calls; free-tier bursts of 8~10 calls produced 429.
RATE_LIMIT_SLEEP = 3


def _need(mod: str):
    try:
        return importlib.import_module(mod)
    except ImportError:
        pkg = DEPS.get(mod, mod)
        print(
            f"Error: '{pkg}' is missing; this step cannot continue.\n"
            f"  Install: uv pip install -r requirements.txt   (deck skill directory)\n"
            f"  Or: uv pip install {pkg}",
            file=sys.stderr,
        )
        raise SystemExit(2)


def load_call_gemini():
    """Load call_gemini from the sibling image skill by relative path.

    Assume only sibling placement, without absolute paths or plugin-root
    environment variables that vary by installation.
    """
    target = Path(__file__).resolve().parents[2] / "image" / "scripts" / "nanobanana.py"
    if not target.is_file():
        print(
            "Error: cannot find nanobanana.py in the image skill.\n"
            f"  Checked location: {target}\n"
            "  Upscaling belongs to image. Use --skip-upscale to run without that skill.",
            file=sys.stderr,
        )
        raise SystemExit(2)
    spec = importlib.util.spec_from_file_location("_deck_nanobanana", target)
    mod = importlib.util.module_from_spec(spec)
    try:
        spec.loader.exec_module(mod)
    except ImportError:
        print(
            "Error: missing image-skill dependencies; upscaling is unavailable.\n"
            "  Install: uv pip install -r requirements.txt\n"
            "  Or use --skip-upscale to retain original resolution.",
            file=sys.stderr,
        )
        raise SystemExit(2)
    fn = getattr(mod, "call_gemini", None)
    if fn is None:
        print("Error: skills/image/scripts/nanobanana.py has no call_gemini function.", file=sys.stderr)
        raise SystemExit(2)
    return fn


# ── Step 1: extract ─────────────────────────────────────────────────────────────


def extract(pdf_path: str) -> list[dict]:
    """Extract original embedded PDF images without rerendering them.

    Rendering at 300 DPI loses original sharpness. Fall back to 2x rendering
    only for pages without embedded images.
    """
    fitz = _need("fitz")
    _need("PIL")
    from PIL import Image

    EXTRACTED_DIR.mkdir(parents=True, exist_ok=True)
    doc = fitz.open(pdf_path)
    pages = []
    for i in range(len(doc)):
        page = doc[i]
        src_path = None
        for info in page.get_images(full=True):
            try:
                base = doc.extract_image(info[0])
            except Exception:
                continue
            if base:
                ext = base.get("ext", "png")
                path = str(EXTRACTED_DIR / f"page_{i + 1:02d}.{ext}")
                Path(path).write_bytes(base["image"])
                src_path = path
                break
        if not src_path:
            pix = page.get_pixmap(matrix=fitz.Matrix(2, 2))
            src_path = str(EXTRACTED_DIR / f"page_{i + 1:02d}_render.png")
            pix.save(src_path)
        img = Image.open(src_path)
        pages.append({"page_num": i, "source_image": src_path,
                      "width": img.width, "height": img.height,
                      "status": "extracted"})
        print(f"  Page {i + 1}: {img.width}x{img.height} -> {src_path}")
    doc.close()
    return pages


# ── Step 2: upscale ─────────────────────────────────────────────────────────


def upscale(pages: list[dict]) -> list[dict]:
    call_gemini = load_call_gemini()
    UPSCALED_DIR.mkdir(parents=True, exist_ok=True)
    for page in pages:
        if page["status"] == "upscaled":
            print(f"  Page {page['page_num'] + 1}: already complete; skipped")
            continue
        src = page["source_image"]
        dst = str(UPSCALED_DIR / f"page_{page['page_num'] + 1:02d}_4k.png")
        print(f"  Page {page['page_num'] + 1}: upscaling...", end=" ", flush=True)
        try:
            call_gemini(text=UPSCALE_PROMPT, image_path=src, generate=True,
                        image_size="4K", aspect_ratio="16:9", output=dst)
            page["upscaled_path"] = dst
            page["status"] = "upscaled"
            print("OK")
        except Exception as e:                      # noqa: BLE001
            # Fall back to the original for failed pages; retain the rest of the deck.
            print(f"failed ({e})")
            page["upscaled_path"] = src
            page["status"] = "upscale_failed"
        time.sleep(RATE_LIMIT_SLEEP)
    return pages


# ── Step 3: assemble ─────────────────────────────────────────────────────────────


def _fit(img_path: str):
    from PIL import Image

    img = Image.open(img_path)
    ratio = img.width / img.height
    if ratio > SLIDE_W_EMU / SLIDE_H_EMU:
        w, h = SLIDE_W_EMU, int(SLIDE_W_EMU / ratio)
        left, top = 0, (SLIDE_H_EMU - h) // 2
    else:
        h, w = SLIDE_H_EMU, int(SLIDE_H_EMU * ratio)
        left, top = (SLIDE_W_EMU - w) // 2, 0
    return left, top, w, h


def assemble(pages: list[dict], output_base: str) -> str:
    _need("pptx")
    _need("PIL")
    from pptx import Presentation
    from pptx.util import Emu

    prs = Presentation()
    prs.slide_width = Emu(SLIDE_W_EMU)
    prs.slide_height = Emu(SLIDE_H_EMU)
    blank = prs.slide_layouts[6]
    for page in pages:
        slide = prs.slides.add_slide(blank)
        img_path = page.get("upscaled_path", page["source_image"])
        left, top, w, h = _fit(img_path)
        slide.shapes.add_picture(img_path, Emu(left), Emu(top), Emu(w), Emu(h))
        print(f"  Slide {page['page_num'] + 1}: {img_path}")

    pptx_path = output_base if output_base.endswith(".pptx") else output_base + ".pptx"
    Path(pptx_path).parent.mkdir(parents=True, exist_ok=True)
    prs.save(pptx_path)
    print(f"  PPTX: {pptx_path} ({Path(pptx_path).stat().st_size / (1024 * 1024):.1f}MB)")

    _need("fpdf")
    from fpdf import FPDF
    from PIL import Image

    pdf_path = pptx_path.replace(".pptx", ".pdf")
    page_w = SLIDE_W_EMU / 914400 * 72
    page_h = SLIDE_H_EMU / 914400 * 72
    pdf = FPDF(orientation="L", unit="pt", format=(page_h, page_w))
    pdf.set_auto_page_break(False)
    for page in pages:
        pdf.add_page()
        img_path = page.get("upscaled_path", page["source_image"])
        img = Image.open(img_path)
        ratio = img.width / img.height
        if ratio > page_w / page_h:
            w, h = page_w, page_w / ratio
            x, y = 0.0, (page_h - h) / 2
        else:
            h, w = page_h, page_h * ratio
            x, y = (page_w - w) / 2, 0.0
        pdf.image(img_path, x=x, y=y, w=w, h=h)
    pdf.output(pdf_path)
    print(f"  PDF:  {pdf_path} ({Path(pdf_path).stat().st_size / (1024 * 1024):.1f}MB)")
    return pptx_path


# ── State ─────────────────────────────────────────────────────────────────────


def save_state(pdf_path: str, pages: list[dict]):
    STATE_PATH.parent.mkdir(parents=True, exist_ok=True)
    STATE_PATH.write_text(
        json.dumps({"pdf_path": pdf_path, "pages": pages},
                   ensure_ascii=False, indent=2), encoding="utf-8")


def load_state() -> dict | None:
    if STATE_PATH.exists():
        return json.loads(STATE_PATH.read_text(encoding="utf-8"))
    return None


def main():
    parser = argparse.ArgumentParser(
        description="Reassemble a PDF presentation into PPTX + PDF (optional 4K upscaling)")
    parser.add_argument("pdf", help="Input PDF")
    parser.add_argument("--skip-upscale", action="store_true",
                        help="Assemble at original resolution without upscaling (no API key required)")
    parser.add_argument("--resume", action="store_true",
                        help="Resume from workspace/state.json")
    parser.add_argument("-o", "--output", default=str(OUTPUT_DIR / "result.pptx"),
                        help="Output PPTX path (PDF is created alongside it)")
    args = parser.parse_args()

    if not Path(args.pdf).exists():
        print(f"Error: {args.pdf} does not exist", file=sys.stderr)
        return 1

    print("=== PDF -> PPTX ===")
    print(f"Input: {args.pdf}")

    if args.resume and (state := load_state()):
        pages = state["pages"]
        print(f"[extract] resume {len(pages)} pages from state.json")
    else:
        print("[extract] extracting embedded images...")
        pages = extract(args.pdf)
        save_state(args.pdf, pages)

    if not pages:
        print("Error: zero pages", file=sys.stderr)
        return 1

    if args.skip_upscale:
        print("[upscale] skipped; retain original resolution")
        for p in pages:
            p.setdefault("upscaled_path", p["source_image"])
    else:
        print(f"[upscale] converting {len(pages)} pages to 4K...")
        pages = upscale(pages)
        save_state(args.pdf, pages)

    print("[assemble] PPTX + PDF...")
    assemble(pages, args.output)
    print("=== Complete ===")
    return 0


if __name__ == "__main__":
    sys.exit(main())
